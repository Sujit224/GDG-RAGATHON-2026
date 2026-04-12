# # import streamlit as st
# # from extractor import extract_from_text
# # from rag_engine import calculate_readiness_score

# # st.set_page_config(page_title="Placement Predictor", layout="centered")

# # st.title("🚀 Placement Readiness Predictor")

# # # -----------------------------
# # # Option Selection
# # # -----------------------------
# # option = st.radio(
# #     "Choose Input Method",
# #     ["Paste Resume", "Manual Input"]
# # )

# # data = None

# # # -----------------------------
# # # OPTION 1: Resume Input
# # # -----------------------------
# # if option == "Paste Resume":
# #     resume_text = st.text_area("Paste your resume here")

# #     if st.button("Analyze Resume"):
# #         if resume_text.strip() == "":
# #             st.warning("Please paste your resume")
# #         else:
# #             data = extract_from_text(resume_text)

# # # -----------------------------
# # # OPTION 2: Manual Input
# # # -----------------------------
# # elif option == "Manual Input":
# #     cgpa = st.slider("CGPA", 0.0, 10.0, 7.0)
# #     projects = st.number_input("Projects", 0, 10, 1)
# #     internships = st.number_input("Internships", 0, 5, 0)
# #     skills = st.number_input("Skills", 0, 20, 3)
# #     dsa = st.slider("DSA Level", 0, 10, 5)

# #     if st.button("Calculate Score"):
# #         data = {
# #             "cgpa": cgpa,
# #             "projects": projects,
# #             "internships": internships,
# #             "skills": skills,
# #             "dsa": dsa
# #         }

# # # -----------------------------
# # # RESULT SECTION
# # # -----------------------------
# # if data:
# #     score = calculate_readiness_score(data)

# #     st.subheader("📊 Extracted Data")
# #     st.json(data)

# #     st.subheader("🔥 Readiness Score")
# #     st.success(f"{score}/100")

# #     # Feedback
# #     feedback = []

# #     if data["projects"] < 2:
# #         feedback.append("Build more projects")

# #     if data["internships"] == 0:
# #         feedback.append("Try getting an internship")

# #     if data["dsa"] < 5:
# #         feedback.append("Improve DSA")

# #     if feedback:
# #         st.subheader("💡 Suggestions")
# #         for tip in feedback:
# #             st.write(f"- {tip}")

# import streamlit as st
# from extractor import extract_from_text

# # -----------------------------
# # Resume Upload
# # -----------------------------
# uploaded_file = st.file_uploader(
#     "Upload Resume",
#     type=["pdf", "txt", "docx"]
# )

# resume_text = ""

# if uploaded_file:
#     file_type = uploaded_file.name.split(".")[-1]

#     if file_type == "txt":
#         resume_text = str(uploaded_file.read(), "utf-8")

#     elif file_type == "pdf":
#         import PyPDF2
#         pdf_reader = PyPDF2.PdfReader(uploaded_file)
#         for page in pdf_reader.pages:
#             resume_text += page.extract_text()

#     elif file_type == "docx":
#         import docx
#         doc = docx.Document(uploaded_file)
#         for para in doc.paragraphs:
#             resume_text += para.text + "\n"

#     st.success("✅ Resume uploaded successfully!")

# # -----------------------------
# # Analyze Button
# # -----------------------------
# if st.button("Analyze Resume"):
#     if resume_text.strip() == "":
#         st.warning("Please upload or paste resume")
#     else:
#         data = extract_from_text(resume_text)

#         from rag_engine import calculate_readiness_score
#         score = calculate_readiness_score(data)

#         st.subheader("📊 Extracted Data")
#         st.json(data)

#         st.subheader("🔥 Readiness Score")
#         st.success(f"{score}/100")


import streamlit as st
from extractor import extract_from_text
from rag_engine import calculate_readiness_score

st.set_page_config(page_title="Placement Predictor", layout="centered")
st.title("🚀 Placement Readiness Predictor")

# -----------------------------
# Mode Selection
# -----------------------------
mode = st.radio(
    "Choose Input Method",
    ["Resume Upload", "Chat Input"]
)

data = None

# =========================================================
# 📄 RESUME UPLOAD
# =========================================================
if mode == "Resume Upload":
    uploaded_file = st.file_uploader(
        "Upload Resume",
        type=["pdf", "txt", "docx"]
    )

    resume_text = ""

    if uploaded_file:
        file_type = uploaded_file.name.split(".")[-1]

        # TXT
        if file_type == "txt":
            resume_text = str(uploaded_file.read(), "utf-8")

        # PDF
        elif file_type == "pdf":
            import PyPDF2
            pdf_reader = PyPDF2.PdfReader(uploaded_file)
            for page in pdf_reader.pages:
                resume_text += page.extract_text() or ""

        # DOCX
        elif file_type == "docx":
            import docx
            doc = docx.Document(uploaded_file)
            for para in doc.paragraphs:
                resume_text += para.text + "\n"

        st.success("✅ Resume uploaded successfully!")

    if st.button("Analyze Resume"):
        if resume_text.strip() == "":
            st.warning("Please upload a resume")
        else:
            data = extract_from_text(resume_text)

# =========================================================
# 💬 CHAT INPUT
# =========================================================
elif mode == "Chat Input":
    user_input = st.text_area(
        "Describe your profile",
        placeholder="Example: I have 8 CGPA, 2 projects, 1 internship, know Python and DSA"
    )

    if st.button("Analyze Chat"):
        if user_input.strip() == "":
            st.warning("Please enter your details")
        else:
            data = extract_from_text(user_input)

# =========================================================
# 📊 RESULT SECTION
# =========================================================
if data:
    score = calculate_readiness_score(data)

    st.subheader("📊 Extracted Data")
    st.json(data)

    st.subheader("🔥 Readiness Score")
    st.success(f"{score}/100")

    # -----------------------------
    # Suggestions
    # -----------------------------
    st.subheader("💡 Suggestions")

    suggestions = []

    if data["projects"] < 2:
        suggestions.append("Build more projects")

    if data["internships"] == 0:
        suggestions.append("Try getting an internship")

    if data["dsa"] < 5:
        suggestions.append("Improve DSA skills")

    if data["skills"] < 5:
        suggestions.append("Learn more relevant skills")

    if not suggestions:
        st.success("Great profile! Keep going 🚀")
    else:
        for s in suggestions:
            st.write(f"- {s}")